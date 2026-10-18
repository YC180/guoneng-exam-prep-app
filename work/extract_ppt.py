#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""提取 PPT 文本（.ppt 走 PowerPoint COM，.pptx 走 python-pptx），写入 work/text6。

背景：`01面试装着讲义(1).ppt` 是 356 份资料里**唯一**的非 PDF/DOCX/TXT/图片文件，
早期提取脚本与覆盖度审计都把它静默跳过了 —— 正是本项目明令禁止的“静默吞掉”。
"""
import os
import sys

WORK = os.path.dirname(os.path.abspath(__file__))
ROOT = r"D:\DOWNLOADS\26国家能源集团通关资料"
OUT = os.path.join(WORK, "text6")


def extract_ppt_com(path):
    """旧版 .ppt 只有 PowerPoint COM 能可靠读取。"""
    import win32com.client as wc
    app = wc.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(path, WithWindow=False)
        try:
            parts = []
            for i in range(1, pres.Slides.Count + 1):
                parts.append(f"[[SLIDE {i}]]")
                for sh in pres.Slides(i).Shapes:
                    try:
                        if sh.HasTextFrame and sh.TextFrame.HasText:
                            parts.append(sh.TextFrame.TextRange.Text)
                    except Exception:
                        continue
            return "\n".join(parts), pres.Slides.Count
        finally:
            pres.Close()
    finally:
        app.Quit()


def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[[SLIDE {i}]]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts), len(prs.slides)


def main():
    targets = []
    for dp, _dn, fs in os.walk(ROOT):
        for fn in fs:
            if os.path.splitext(fn)[1].lower() in {".ppt", ".pptx"}:
                targets.append(os.path.join(dp, fn))
    print(f"发现 PPT 文件 {len(targets)} 个")
    ok = 0
    for src in targets:
        rel = os.path.relpath(src, ROOT)
        out_path = os.path.join(OUT, os.path.splitext(rel)[0] + ".txt")
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        try:
            if src.lower().endswith(".pptx"):
                body, n = extract_pptx(src)
                how = "python-pptx"
            else:
                body, n = extract_ppt_com(src)
                how = "PowerPoint COM"
        except Exception as exc:
            print(f"ERR {os.path.basename(rel)[:50]} :: {type(exc).__name__}: {exc}")
            continue
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(f"# SOURCE: {src}\n# EXTRACT: {how}\n\n{body}\n")
        print(f"OK  {n:3d} 页 {len(body):6d} 字符  {os.path.basename(rel)[:50]}")
        ok += 1
    print(f"\nPPT 提取成功 {ok}/{len(targets)}")
    return 0 if ok == len(targets) else 2


if __name__ == "__main__":
    sys.exit(main())
