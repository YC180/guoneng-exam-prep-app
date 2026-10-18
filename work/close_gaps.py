#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""补齐覆盖度审计发现的未读文件，并核实 7、与 9、目录是否为真重复。

- A: 用 MD5 核对 7、与 9、两目录文件是否逐一相同（相同则无需重复读取）
- B: 把审计中“无任何产出”的文件逐个提取到 work/text5（命名保留扩展名，避免同名覆盖）
"""
import hashlib
import json
import os
import sys
import time
import traceback

WORK = os.path.dirname(os.path.abspath(__file__))
ROOT = r"D:\DOWNLOADS\26国家能源集团通关资料"
OUT = os.path.join(WORK, "text5")


def md5(p, chunk=1 << 20):
    h = hashlib.md5()
    with open(p, "rb") as fh:
        while True:
            b = fh.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def dup_check():
    d7 = os.path.join(ROOT, "7、【EPI通用能力练习题库】国家能源集团招聘")
    d9 = os.path.join(ROOT, "9、国家能源集团招聘EPI通用能力练习题库")
    if not (os.path.isdir(d7) and os.path.isdir(d9)):
        return None
    l7 = {}
    for dp, _dn, fns in os.walk(d7):
        for fn in fns:
            l7[os.path.relpath(os.path.join(dp, fn), d7)] = os.path.join(dp, fn)
    same = diff = missing = 0
    diffs = []
    for rel, p7 in sorted(l7.items()):
        p9 = os.path.join(d9, rel)
        if not os.path.exists(p9):
            missing += 1
            diffs.append(f"缺失于9、: {rel}")
            continue
        if os.path.getsize(p7) == os.path.getsize(p9) and md5(p7) == md5(p9):
            same += 1
        else:
            diff += 1
            diffs.append(f"内容不同: {rel}")
    print(f"[A] 7、共 {len(l7)} 文件 → 与 9、逐一 MD5 相同 {same}，不同 {diff}，9、缺失 {missing}")
    for d in diffs[:10]:
        print("     ", d)
    return {"same": same, "diff": diff, "missing": missing, "details": diffs}


def safe_out(rel):
    return os.path.join(OUT, rel.replace("\\", "__") + ".txt")


def extract(src):
    ext = os.path.splitext(src)[1].lower()
    if ext == ".docx":
        import docx
        d = docx.Document(src)
        parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
        for tbl in d.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "python-docx"
    if ext == ".doc":
        import pythoncom
        import win32com.client as wc
        pythoncom.CoInitialize()
        app = doc = None
        try:
            app = wc.Dispatch("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            doc = app.Documents.Open(src, ReadOnly=True, AddToRecentFiles=False)
            return doc.Content.Text, "word-com"
        finally:
            try:
                if doc is not None:
                    doc.Close(False)
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            pythoncom.CoUninitialize()
    if ext in (".ppt", ".pptx"):
        import pythoncom
        import win32com.client as wc
        pythoncom.CoInitialize()
        app = pres = None
        try:
            app = wc.Dispatch("PowerPoint.Application")
            pres = app.Presentations.Open(src, WithWindow=False, ReadOnly=True)
            buf = []
            for i in range(1, pres.Slides.Count + 1):
                sl = pres.Slides(i)
                for j in range(1, sl.Shapes.Count + 1):
                    try:
                        if sl.Shapes(j).HasTextFrame:
                            t = sl.Shapes(j).TextFrame.TextRange.Text
                            if t and t.strip():
                                buf.append(f"[幻灯片 {i}] {t.strip()}")
                    except Exception:
                        continue
            return "\n".join(buf), "powerpoint-com"
        finally:
            try:
                if pres is not None:
                    pres.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            pythoncom.CoUninitialize()
    if ext == ".pdf":
        import fitz
        doc = fitz.open(src)
        txt = "\n".join(doc.load_page(i).get_text() for i in range(len(doc)))
        doc.close()
        return txt, "pymupdf-text"
    if ext in (".txt", ".md", ".csv", ".json"):
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                return open(src, encoding=enc).read(), f"plain({enc})"
            except (UnicodeDecodeError, UnicodeError):
                continue
        return open(src, encoding="utf-8", errors="replace").read(), "plain(replace)"
    raise ValueError(f"不支持的类型 {ext}")


def main():
    os.makedirs(OUT, exist_ok=True)
    print("[A] 核实 7、/9、是否重复")
    dup = dup_check()

    aud = json.load(open(os.path.join(WORK, "bank", "coverage_audit.json"), encoding="utf-8"))
    unread = [u["file"] for u in aud["unread"]]
    dup_prefix = "7、【EPI通用能力练习题库】国家能源集团招聘\\"
    skip_dup = bool(dup and dup["diff"] == 0 and dup["missing"] == 0)
    todo = []
    skipped = 0
    for rel in unread:
        if skip_dup and rel.startswith(dup_prefix):
            skipped += 1
            continue
        todo.append(rel)
    print(f"\n[B] 未读 {len(unread)} 个：跳过已证实重复 {skipped} 个，待提取 {len(todo)} 个")

    ok = err = 0
    for rel in todo:
        src = os.path.join(ROOT, rel)
        if not os.path.exists(src):
            print(f"  MISSING {rel}")
            continue
        op = safe_out(rel)
        if os.path.exists(op) and os.path.getsize(op) > 300:
            ok += 1
            continue
        t0 = time.time()
        try:
            body, how = extract(src)
            with open(op, "w", encoding="utf-8") as fh:
                fh.write(f"# SOURCE: {src}\n# HOW: {how}\n\n{body}\n")
            ok += 1
            print(f"  OK {len(body):7d}c {time.time()-t0:5.1f}s [{how:16s}] {os.path.basename(rel)[:54]}")
        except Exception as exc:
            err += 1
            print(f"  ERR {os.path.basename(rel)[:54]} :: {type(exc).__name__}: {exc}")
            traceback.print_exc()
    print(f"\n[B] 完成 ok={ok} err={err}")
    json.dump({"dup7vs9": dup, "skippedAsDuplicate": skipped, "extracted": ok, "errors": err},
              open(os.path.join(WORK, "bank", "gap_closure.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    return 0


if __name__ == "__main__":
    sys.exit(main())
