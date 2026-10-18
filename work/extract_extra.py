#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""补救提取：处理覆盖度审计发现的“从未被打开”文件。

- .docx → python-docx
- .doc  → Word COM（pywin32）
- .pdf  → PyMuPDF 取文字层；若无文字层则 RapidOCR 识别

输出到 work/text5，带 `# SOURCE:` 头，便于覆盖度审计回填映射。
"""
import os
import sys
import time
import traceback

WORK = os.path.dirname(os.path.abspath(__file__))
ROOT = r"D:\DOWNLOADS\26国家能源集团通关资料"
OUT = os.path.join(WORK, "text5")

FILES = [
    r"1、2025国家能源集团招聘考试知识点讲义和企业文化\国家能源集团企业文化.doc",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\2025年中央一号文件\2025年中央一号文件-资料包\【模拟试题】2025年中央一号文件-答案解析.docx",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\2025年中央一号文件\2025年中央一号文件-资料包\【模拟试题】2025年中央一号文件-题目.docx",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\2025年中央一号文件\2025年中央一号文件-资料包\【考点集萃】2025年中央一号文件.docx",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\2025年中央一号文件\2025年中央一号文件-资料包\【重点标注】2025年中央一号文件.docx",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\重要讲话包\❤【月更】24年-25年重要讲话考点预测400题（更新至25年7月底）.pdf",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\重要讲话包\❤【月更】24年-25年重要讲话考点预测400题（更新至25年7月底）-纯题本.pdf",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\【重要会议】2025重要时政+文章考点\重要讲话包\考点速记  2025.9.29中共中央政治局会议.pdf",
    r"8、时政【全年持续更新】\【全国时政】全国时政热点（已更4月）\赠品：学习资料知识库（考公考编考研教招教资龙虾AI等资料）\学习资料知识库（先保存防失效）(1).docx",
]


def out_path_for(rel):
    safe = rel.replace("\\", "__")
    return os.path.join(OUT, safe + ".txt")


def from_docx(src):
    import docx
    d = docx.Document(src)
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def from_doc_com(src):
    import pythoncom  # noqa: F401
    import win32com.client as wc
    pythoncom.CoInitialize()
    app = None
    doc = None
    try:
        app = wc.Dispatch("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0
        doc = app.Documents.Open(src, ReadOnly=True, AddToRecentFiles=False)
        return doc.Content.Text
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


_engine = None


def ocr_pdf(src, dpi=200):
    global _engine
    import fitz
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    if _engine is None:
        _engine = RapidOCR()
    doc = fitz.open(src)
    chunks = []
    empty = 0
    for pno in range(len(doc)):
        pix = doc.load_page(pno).get_pixmap(dpi=dpi)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        if pix.n == 4:
            img = img[:, :, :3]
        res, _el = _engine(img)
        texts = [r[1] for r in res] if res else []
        if not texts:
            empty += 1
        chunks.append(f"[[PAGE {pno + 1}]]\n" + "\n".join(texts))
    doc.close()
    return "\n\n".join(chunks), len(doc), empty


def from_pdf(src):
    import fitz
    doc = fitz.open(src)
    txt = "\n".join(doc.load_page(i).get_text() for i in range(len(doc)))
    doc.close()
    if len(txt.strip()) >= 200:
        return txt, "text-layer", 0
    body, pages, empty = ocr_pdf(src)
    return body, f"ocr(dpi=200,{pages}p,empty={empty})", empty


def main():
    os.makedirs(OUT, exist_ok=True)
    ok = 0
    for rel in FILES:
        src = os.path.join(ROOT, rel)
        if not os.path.exists(src):
            print(f"MISSING  {rel}")
            continue
        ext = os.path.splitext(src)[1].lower()
        t0 = time.time()
        try:
            if ext == ".docx":
                body, how = from_docx(src), "python-docx"
            elif ext == ".doc":
                body, how = from_doc_com(src), "word-com"
            elif ext == ".pdf":
                body, how, _e = from_pdf(src)
            else:
                print(f"SKIP ext {ext}  {rel}")
                continue
            op = out_path_for(rel)
            with open(op, "w", encoding="utf-8") as fh:
                fh.write(f"# SOURCE: {src}\n# HOW: {how}\n\n{body}\n")
            ok += 1
            print(f"OK   {len(body):7d}c {time.time()-t0:6.1f}s [{how:24s}] {os.path.basename(rel)[:56]}")
        except Exception as exc:
            print(f"ERR  {os.path.basename(rel)[:56]} :: {type(exc).__name__}: {exc}")
            traceback.print_exc()
    print(f"\n完成 {ok}/{len(FILES)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
